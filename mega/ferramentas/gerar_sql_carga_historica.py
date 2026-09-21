# -*- coding: utf-8 -*-
"""Gera script SQL para carga historica das planilhas do Approvo.

Le todas as planilhas da pasta APPROVO, aplica as regras de negocio:
- Conversao de datas para padrao ISO YYYY-MM-DD
- Separacao de Data/Hora em data_aprovacao e hora_aprovacao
- Extracao da obra a partir da filial
- Limpeza e parsing de valores monetarios e inteiros
- Deduplicacao inteligente
- Gera comandos INSERT com ON CONFLICT (idempotente)
"""
import datetime as dt
import json
import math
import os
import re
import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
import approvo_carregar

PASTA_APPROVO = Path(r"C:\Users\HomePC\OneDrive - Piemonte Construtora\PLANEJAMENTO PIEMONTE\3.METAS SETOR\2026\Micrometas setor 2026\Acompanhamento de suprimentos\APPROVO")
ARQUIVO_SAIDA_SQL = Path(__file__).resolve().parent.parent / "dados" / "carga_historica_approvo.sql"
ARQUIVO_SAIDA_SQL.parent.mkdir(parents=True, exist_ok=True)


def sql_escape(v):
    if v is None:
        return "NULL"
    if isinstance(v, (int, float)):
        return str(v)
    s = str(v).replace("'", "''")
    return f"'{s}'"


def processar_historico():
    print(f"Buscando arquivos na pasta: {PASTA_APPROVO}")
    if not PASTA_APPROVO.exists():
        print(f"ERRO: Pasta {PASTA_APPROVO} nao encontrada!")
        return

    arquivos_docs = sorted(list(PASTA_APPROVO.glob("*Documentos*.xlsx")))
    arquivos_ocorr = sorted(list(PASTA_APPROVO.glob("*Ocorr*.xlsx")))

    print(f"Encontrados {len(arquivos_docs)} arquivos de Documentos e {len(arquivos_ocorr)} de Ocorrencias.")

    # 1. Processar Documentos
    docs_map = {}  # chave: (filial, tipo_documento, numero) -> dados
    total_linhas_docs = 0

    print("\n--- Processando Documentos ---")
    for f in arquivos_docs:
        print(f"Lendo: {f.name}...")
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows or len(rows) < 2:
            continue

        headers = [str(c).strip() if c is not None else "" for c in rows[0]]
        col_idx = {h: i for i, h in enumerate(headers)}

        linhas_arquivo = 0
        for r in rows[1:]:
            total_linhas_docs += 1
            linhas_arquivo += 1
            numero = approvo_carregar._parse_int(r[col_idx.get('Número', 0)])
            if numero is None:
                continue

            tipo_doc = str(r[col_idx.get('Tipo de documento', 1)] or '').strip()
            data_doc = approvo_carregar._parse_date(r[col_idx.get('Data', 2)])
            status = str(r[col_idx.get('Status', 3)] or '').strip()
            valor = approvo_carregar._parse_float(r[col_idx.get('Valor (R$)', 4)])
            filial = str(r[col_idx.get('Filial', 5)] or '').strip()
            agente = str(r[col_idx.get('Agente', 6)] or '').strip() if r[col_idx.get('Agente', 6)] is not None else None
            solicitante = str(r[col_idx.get('Solicitante', 7)] or '').strip() if r[col_idx.get('Solicitante', 7)] is not None else None
            data_envio = approvo_carregar._parse_date(r[col_idx.get('Data de envio para aprovação', 8)])
            classe_fin = str(r[col_idx.get('Classe financeira', 9)] or '').strip() if r[col_idx.get('Classe financeira', 9)] is not None else None
            centro_custo = str(r[col_idx.get('Centro de custo', 10)] or '').strip() if r[col_idx.get('Centro de custo', 10)] is not None else None
            projeto = str(r[col_idx.get('Projeto', 11)] or '').strip() if r[col_idx.get('Projeto', 11)] is not None else None
            regra_aprov = str(r[col_idx.get('Regra de aprovação', 12)] or '').strip() if r[col_idx.get('Regra de aprovação', 12)] is not None else None
            obra = approvo_carregar._extrair_obra(filial)

            raw_data = json.dumps({
                headers[i]: r[i] for i in range(len(headers)) if i < len(r)
            }, default=str, ensure_ascii=False)

            chave = (filial, tipo_doc, numero)
            # Mantem o mais recente se ja existir
            docs_map[chave] = (
                obra, numero, tipo_doc, data_doc, status, valor,
                filial, agente, solicitante, data_envio, classe_fin,
                centro_custo, projeto, regra_aprov, "2026-09-21", raw_data
            )
        print(f"  -> {linhas_arquivo} linhas lidas.")

    print(f"\nTotal lido em Documentos: {total_linhas_docs} linhas.")
    print(f"Total de Documentos unicos apos deduplicacao: {len(docs_map)}")

    # 2. Processar Ocorrências
    ocorr_set = {}  # chave unica -> dados
    total_linhas_ocorr = 0

    print("\n--- Processando Ocorrencias ---")
    for f in arquivos_ocorr:
        print(f"Lendo: {f.name}...")
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows or len(rows) < 2:
            continue

        headers = [str(c).strip() if c is not None else "" for c in rows[0]]
        col_idx = {h: i for i, h in enumerate(headers)}

        linhas_arquivo = 0
        for r in rows[1:]:
            total_linhas_ocorr += 1
            linhas_arquivo += 1
            numero_doc = approvo_carregar._parse_int(r[col_idx.get('Número do documento', 0)])
            if numero_doc is None:
                continue

            acao = str(r[col_idx.get('Ação', 1)] or '').strip()
            data_hora_raw = r[col_idx.get('Data/Hora', 2)]
            data_aprov, hora_aprov, ts, data_hora_texto = approvo_carregar._parse_data_hora(data_hora_raw)
            aprovador = str(r[col_idx.get('Aprovador', 3)] or '').strip()
            motivo = str(r[col_idx.get('Motivo da operação', 4)] or '').strip() if r[col_idx.get('Motivo da operação', 4)] is not None else None
            tipo_doc = str(r[col_idx.get('Tipo de documento', 5)] or '').strip()
            data_doc = approvo_carregar._parse_date(r[col_idx.get('Data do documento', 6)])
            valor = approvo_carregar._parse_float(r[col_idx.get('Valor (R$)', 7)])
            filial = str(r[col_idx.get('Filial', 8)] or '').strip()
            agente = str(r[col_idx.get('Agente', 9)] or '').strip() if r[col_idx.get('Agente', 9)] is not None else None
            solicitante = str(r[col_idx.get('Solicitante', 10)] or '').strip() if r[col_idx.get('Solicitante', 10)] is not None else None
            obra = approvo_carregar._extrair_obra(filial)

            chave_unica = (filial, tipo_doc, numero_doc, data_aprov, hora_aprov, aprovador, acao)
            if chave_unica in ocorr_set:
                continue

            raw_data = json.dumps({
                headers[i]: r[i] for i in range(len(headers)) if i < len(r)
            }, default=str, ensure_ascii=False)

            ocorr_set[chave_unica] = (
                obra, numero_doc, acao, data_aprov, hora_aprov,
                ts, data_hora_texto, aprovador, motivo, tipo_doc,
                data_doc, valor, filial, agente, solicitante, "2026-09-21", raw_data
            )
        print(f"  -> {linhas_arquivo} linhas lidas.")

    print(f"\nTotal lido em Ocorrencias: {total_linhas_ocorr} linhas.")
    print(f"Total de Ocorrencias unicas apos deduplicacao: {len(ocorr_set)}")

    # 3. Gerar Arquivo SQL
    print(f"\nEscrevendo script SQL em: {ARQUIVO_SAIDA_SQL}...")
    with open(ARQUIVO_SAIDA_SQL, "w", encoding="utf-8") as sql_file:
        sql_file.write("-- CARGA HISTORICA APPROVO MEGA ERP\n")
        sql_file.write("-- Gerado automaticamente para backfill completo\n")
        sql_file.write("BEGIN;\n\n")

        # Insercoes de Documentos em blocos de 500
        docs_list = list(docs_map.values())
        batch_size = 500
        for i in range(0, len(docs_list), batch_size):
            batch = docs_list[i : i + batch_size]
            sql_file.write("INSERT INTO mega.approvo_documentos (\n")
            sql_file.write("  obra, numero, tipo_documento, data_documento, status, valor,\n")
            sql_file.write("  filial, agente, solicitante, data_envio_aprovacao, classe_financeira,\n")
            sql_file.write("  centro_custo, projeto, regra_aprovacao, data_extracao, raw_data\n")
            sql_file.write(") VALUES\n")

            values_lines = []
            for d in batch:
                v_str = f"  ({sql_escape(d[0])}, {d[1]}, {sql_escape(d[2])}, {sql_escape(d[3])}, {sql_escape(d[4])}, {sql_escape(d[5])}, {sql_escape(d[6])}, {sql_escape(d[7])}, {sql_escape(d[8])}, {sql_escape(d[9])}, {sql_escape(d[10])}, {sql_escape(d[11])}, {sql_escape(d[12])}, {sql_escape(d[13])}, {sql_escape(d[14])}, {sql_escape(d[15])}::jsonb)"
                values_lines.append(v_str)

            sql_file.write(",\n".join(values_lines))
            sql_file.write("\nON CONFLICT (filial, tipo_documento, numero) DO UPDATE SET\n")
            sql_file.write("  obra = EXCLUDED.obra,\n")
            sql_file.write("  data_documento = EXCLUDED.data_documento,\n")
            sql_file.write("  status = EXCLUDED.status,\n")
            sql_file.write("  valor = EXCLUDED.valor,\n")
            sql_file.write("  agente = EXCLUDED.agente,\n")
            sql_file.write("  solicitante = EXCLUDED.solicitante,\n")
            sql_file.write("  data_envio_aprovacao = EXCLUDED.data_envio_aprovacao,\n")
            sql_file.write("  classe_financeira = EXCLUDED.classe_financeira,\n")
            sql_file.write("  centro_custo = EXCLUDED.centro_custo,\n")
            sql_file.write("  projeto = EXCLUDED.projeto,\n")
            sql_file.write("  regra_aprovacao = EXCLUDED.regra_aprovacao,\n")
            sql_file.write("  data_extracao = EXCLUDED.data_extracao,\n")
            sql_file.write("  raw_data = EXCLUDED.raw_data;\n\n")

        # Insercoes de Ocorrencias em blocos de 500
        ocorr_list = list(ocorr_set.values())
        for i in range(0, len(ocorr_list), batch_size):
            batch = ocorr_list[i : i + batch_size]
            sql_file.write("INSERT INTO mega.approvo_ocorrencias (\n")
            sql_file.write("  obra, numero_documento, acao, data_aprovacao, hora_aprovacao,\n")
            sql_file.write("  data_hora, data_hora_texto, aprovador, motivo_operacao, tipo_documento,\n")
            sql_file.write("  data_documento, valor, filial, agente, solicitante, data_extracao, raw_data\n")
            sql_file.write(") VALUES\n")

            values_lines = []
            for o in batch:
                v_str = f"  ({sql_escape(o[0])}, {o[1]}, {sql_escape(o[2])}, {sql_escape(o[3])}, {sql_escape(o[4])}, {sql_escape(o[5])}, {sql_escape(o[6])}, {sql_escape(o[7])}, {sql_escape(o[8])}, {sql_escape(o[9])}, {sql_escape(o[10])}, {sql_escape(o[11])}, {sql_escape(o[12])}, {sql_escape(o[13])}, {sql_escape(o[14])}, {sql_escape(o[15])}, {sql_escape(o[16])}::jsonb)"
                values_lines.append(v_str)

            sql_file.write(",\n".join(values_lines))
            sql_file.write("\nON CONFLICT (filial, tipo_documento, numero_documento, data_aprovacao, hora_aprovacao, aprovador, acao) DO NOTHING;\n\n")

        sql_file.write("COMMIT;\n")

    tamanho = ARQUIVO_SAIDA_SQL.stat().st_size / (1024 * 1024)
    print(f"\nSUCESSO! Arquivo SQL gerado: {ARQUIVO_SAIDA_SQL} ({tamanho:.2f} MB)")


if __name__ == "__main__":
    processar_historico()
