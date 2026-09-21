# -*- coding: utf-8 -*-
"""Gera script SQL para carga historica a partir do arquivo consolidado POTY.

Arquivo: DATA INPUT MEGA E APROVAÇÕES.xlsx
Abas:
- DOCUMENTOS APROVADOS (108.367 linhas -> ~98.754 documentos unicos)
- OCORRENCIAS DE APROVAÇÃO (164.872 linhas -> ~161.192 ocorrencias unicas)
"""
import datetime as dt
import gzip
import json
import shutil
import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
import approvo_carregar

ARQUIVO_POTY = Path(r"C:\Users\HomePC\OneDrive - Piemonte Construtora\PLANEJAMENTO PIEMONTE\3.METAS SETOR\2026\Micrometas setor 2026\Acompanhamento de suprimentos\POTY\DATA INPUT MEGA E APROVAÇÕES.xlsx")
SQL_OUT = Path(__file__).resolve().parent.parent / "dados" / "carga_historica_approvo_master.sql"
SQL_GZ_OUT = Path(__file__).resolve().parent.parent / "dados" / "carga_historica_approvo_master.sql.gz"
SQL_OUT.parent.mkdir(parents=True, exist_ok=True)


def sql_escape(v):
    if v is None:
        return "NULL"
    if isinstance(v, (int, float)):
        return str(v)
    s = str(v).replace("'", "''")
    return f"'{s}'"


def gerar_sql():
    print(f"Lendo arquivo master: {ARQUIVO_POTY}...")
    wb = openpyxl.load_workbook(ARQUIVO_POTY, read_only=True, data_only=True)

    # 1. DOCUMENTOS
    sheet_doc_name = [s for s in wb.sheetnames if 'DOC' in s][0]
    ws_doc = wb[sheet_doc_name]
    print(f"\n--- Processando aba: {sheet_doc_name} ---")
    headers_doc = None
    docs_map = {}
    total_linhas_doc = 0

    for row in ws_doc.iter_rows(values_only=True):
        if headers_doc is None:
            headers_doc = [str(c).strip() if c is not None else '' for c in row]
            col_doc = {h: i for i, h in enumerate(headers_doc)}
            continue
        if not any(row):
            continue
        total_linhas_doc += 1
        numero = approvo_carregar._parse_int(row[col_doc.get('Número', 0)])
        if numero is None:
            continue

        tipo_doc = str(row[col_doc.get('Tipo de documento', 1)] or '').strip()
        data_doc = approvo_carregar._parse_date(row[col_doc.get('Data', 2)])
        status = str(row[col_doc.get('Status', 3)] or '').strip()
        valor = approvo_carregar._parse_float(row[col_doc.get('Valor (R$)', 4)])
        filial = str(row[col_doc.get('Filial', 5)] or '').strip()
        agente = str(row[col_doc.get('Agente', 6)] or '').strip() if row[col_doc.get('Agente', 6)] is not None else None
        solicitante = str(row[col_doc.get('Solicitante', 7)] or '').strip() if row[col_doc.get('Solicitante', 7)] is not None else None
        data_envio = approvo_carregar._parse_date(row[col_doc.get('Data de envio para aprovação', 8)])
        classe_fin = str(row[col_doc.get('Classe financeira', 9)] or '').strip() if row[col_doc.get('Classe financeira', 9)] is not None else None
        centro_custo = str(row[col_doc.get('Centro de custo', 10)] or '').strip() if row[col_doc.get('Centro de custo', 10)] is not None else None
        projeto = str(row[col_doc.get('Projeto', 11)] or '').strip() if row[col_doc.get('Projeto', 11)] is not None else None
        regra_aprov = str(row[col_doc.get('Regra de aprovação', 12)] or '').strip() if row[col_doc.get('Regra de aprovação', 12)] is not None else None
        obra = approvo_carregar._extrair_obra(filial)

        raw_data = json.dumps({
            headers_doc[i]: row[i] for i in range(len(headers_doc)) if i < len(row)
        }, default=str, ensure_ascii=False)

        chave = (filial, tipo_doc, numero)
        docs_map[chave] = (
            obra, numero, tipo_doc, data_doc, status, valor,
            filial, agente, solicitante, data_envio, classe_fin,
            centro_custo, projeto, regra_aprov, "2026-09-21", raw_data
        )

    print(f"Total lido em Documentos: {total_linhas_doc} linhas.")
    print(f"Total unico de Documentos: {len(docs_map)}")

    # 2. OCORRÊNCIAS
    sheet_ocorr_name = [s for s in wb.sheetnames if 'OCORR' in s][0]
    ws_ocorr = wb[sheet_ocorr_name]
    print(f"\n--- Processando aba: {sheet_ocorr_name} ---")
    headers_ocorr = None
    ocorr_map = {}
    total_linhas_ocorr = 0

    for row in ws_ocorr.iter_rows(values_only=True):
        if headers_ocorr is None:
            headers_ocorr = [str(c).strip() if c is not None else '' for c in row]
            col_ocorr = {h: i for i, h in enumerate(headers_ocorr)}
            continue
        if not any(row):
            continue
        total_linhas_ocorr += 1
        num_doc = approvo_carregar._parse_int(row[col_ocorr.get('Número do documento', 0)])
        if num_doc is None:
            continue

        acao = str(row[col_ocorr.get('Ação', 1)] or '').strip()
        data_hora_raw = row[col_ocorr.get('Data/Hora', 2)]
        data_aprov, hora_aprov, ts, data_hora_texto = approvo_carregar._parse_data_hora(data_hora_raw)
        aprovador = str(row[col_ocorr.get('Aprovador', 3)] or '').strip()
        motivo = str(row[col_ocorr.get('Motivo da operação', 4)] or '').strip() if row[col_ocorr.get('Motivo da operação', 4)] is not None else None
        tipo_doc = str(row[col_ocorr.get('Tipo de documento', 5)] or '').strip()
        data_doc = approvo_carregar._parse_date(row[col_ocorr.get('Data do documento', 6)])
        valor = approvo_carregar._parse_float(row[col_ocorr.get('Valor (R$)', 7)])
        filial = str(row[col_ocorr.get('Filial', 8)] or '').strip()
        agente = str(row[col_ocorr.get('Agente', 9)] or '').strip() if row[col_ocorr.get('Agente', 9)] is not None else None
        solicitante = str(row[col_ocorr.get('Solicitante', 10)] or '').strip() if row[col_ocorr.get('Solicitante', 10)] is not None else None
        obra = approvo_carregar._extrair_obra(filial)

        chave_unica = (filial, tipo_doc, num_doc, data_aprov, hora_aprov, aprovador, acao)
        if chave_unica in ocorr_map:
            continue

        raw_data = json.dumps({
            headers_ocorr[i]: row[i] for i in range(len(headers_ocorr)) if i < len(row)
        }, default=str, ensure_ascii=False)

        ocorr_map[chave_unica] = (
            obra, num_doc, acao, data_aprov, hora_aprov,
            ts, data_hora_texto, aprovador, motivo, tipo_doc,
            data_doc, valor, filial, agente, solicitante, "2026-09-21", raw_data
        )

    print(f"Total lido em Ocorrencias: {total_linhas_ocorr} linhas.")
    print(f"Total unico de Ocorrencias: {len(ocorr_map)}")
    wb.close()

    # 3. Gerar Arquivo SQL
    print(f"\nEscrevendo SQL em: {SQL_OUT}...")
    with open(SQL_OUT, "w", encoding="utf-8") as sql_file:
        sql_file.write("-- CARGA HISTORICA MASTER APPROVO MEGA ERP (POTY)\n")
        sql_file.write("-- Fonte: DATA INPUT MEGA E APROVAÇÕES.xlsx\n")
        sql_file.write("BEGIN;\n\n")

        # Insercoes de Documentos em blocos de 500
        docs_list = list(docs_map.values())
        batch_size = 500
        for i in range(0, len(docs_list), batch_size):
            batch = docs_list[i : i + batch_size]
            sql_file.write("INSERT INTO mega.approvo_documentos (\n")
            sql_file.write("  obra, numero, tipo_documento, data_documento, status, valor,\n")
            sql_file.write("  filial, agente, solicitante, data_envio_aprovacao, classe_financeira,\n")
            sql_file.write("  centro_custo, projeto, regra_aprovacao, data_extracao, raw_data\n")
            sql_file.write(") VALUES\n")

            values_lines = []
            for d in batch:
                v_str = f"  ({sql_escape(d[0])}, {d[1]}, {sql_escape(d[2])}, {sql_escape(d[3])}, {sql_escape(d[4])}, {sql_escape(d[5])}, {sql_escape(d[6])}, {sql_escape(d[7])}, {sql_escape(d[8])}, {sql_escape(d[9])}, {sql_escape(d[10])}, {sql_escape(d[11])}, {sql_escape(d[12])}, {sql_escape(d[13])}, {sql_escape(d[14])}, {sql_escape(d[15])}::jsonb)"
                values_lines.append(v_str)

            sql_file.write(",\n".join(values_lines))
            sql_file.write("\nON CONFLICT (filial, tipo_documento, numero) DO UPDATE SET\n")
            sql_file.write("  obra = EXCLUDED.obra,\n")
            sql_file.write("  data_documento = EXCLUDED.data_documento,\n")
            sql_file.write("  status = EXCLUDED.status,\n")
            sql_file.write("  valor = EXCLUDED.valor,\n")
            sql_file.write("  agente = EXCLUDED.agente,\n")
            sql_file.write("  solicitante = EXCLUDED.solicitante,\n")
            sql_file.write("  data_envio_aprovacao = EXCLUDED.data_envio_aprovacao,\n")
            sql_file.write("  classe_financeira = EXCLUDED.classe_financeira,\n")
            sql_file.write("  centro_custo = EXCLUDED.centro_custo,\n")
            sql_file.write("  projeto = EXCLUDED.projeto,\n")
            sql_file.write("  regra_aprovacao = EXCLUDED.regra_aprovacao,\n")
            sql_file.write("  data_extracao = EXCLUDED.data_extracao,\n")
            sql_file.write("  raw_data = EXCLUDED.raw_data;\n\n")

        # Insercoes de Ocorrencias em blocos de 500
        ocorr_list = list(ocorr_map.values())
        for i in range(0, len(ocorr_list), batch_size):
            batch = ocorr_list[i : i + batch_size]
            sql_file.write("INSERT INTO mega.approvo_ocorrencias (\n")
            sql_file.write("  obra, numero_documento, acao, data_aprovacao, hora_aprovacao,\n")
            sql_file.write("  data_hora, data_hora_texto, aprovador, motivo_operacao, tipo_documento,\n")
            sql_file.write("  data_documento, valor, filial, agente, solicitante, data_extracao, raw_data\n")
            sql_file.write(") VALUES\n")

            values_lines = []
            for o in batch:
                v_str = f"  ({sql_escape(o[0])}, {o[1]}, {sql_escape(o[2])}, {sql_escape(o[3])}, {sql_escape(o[4])}, {sql_escape(o[5])}, {sql_escape(o[6])}, {sql_escape(o[7])}, {sql_escape(o[8])}, {sql_escape(o[9])}, {sql_escape(o[10])}, {sql_escape(o[11])}, {sql_escape(o[12])}, {sql_escape(o[13])}, {sql_escape(o[14])}, {sql_escape(o[15])}, {sql_escape(o[16])}::jsonb)"
                values_lines.append(v_str)

            sql_file.write(",\n".join(values_lines))
            sql_file.write("\nON CONFLICT (filial, tipo_documento, numero_documento, data_aprovacao, hora_aprovacao, aprovador, acao) DO NOTHING;\n\n")

        sql_file.write("COMMIT;\n")

    tamanho = SQL_OUT.stat().st_size / (1024 * 1024)
    print(f"Arquivo SQL gerado: {SQL_OUT} ({tamanho:.2f} MB)")

    # Compacta em gzip
    print(f"Compactando para {SQL_GZ_OUT}...")
    with open(SQL_OUT, 'rb') as f_in:
        with gzip.open(SQL_GZ_OUT, 'wb', compresslevel=9) as f_out:
            shutil.copyfileobj(f_in, f_out)
    tamanho_gz = SQL_GZ_OUT.stat().st_size / (1024 * 1024)
    print(f"SUCESSO TOTAL! Arquivo compactado gerado: {SQL_GZ_OUT} ({tamanho_gz:.2f} MB)")


if __name__ == "__main__":
    gerar_sql()
