# -*- coding: utf-8 -*-
"""Carga dos arquivos Excel do Approvo Mega ERP no PostgreSQL (schema mega).

Realiza o parsing dos dados, converte datas (padrao brasileiro DD/MM/AAAA),
separa a coluna Data/Hora em data_aprovacao e hora_aprovacao, extrai a obra a
partir da filial e executa UPSERT transacional.
"""
import datetime as dt
import json
import math
import os
import re
from pathlib import Path
import openpyxl


def _extrair_obra(filial):
    if not filial:
        return None
    m = re.search(r'^\s*(\d+)', str(filial))
    return m.group(1) if m else None


def _parse_date(v):
    if not v:
        return None
    if isinstance(v, (dt.date, dt.datetime)):
        return v.strftime('%Y-%m-%d')
    m = re.search(r'(\d{2})/(\d{2})/(\d{4})', str(v))
    if m:
        return f'{m.group(3)}-{m.group(2)}-{m.group(1)}'
    return None


def _parse_data_hora(v):
    """Extrai (data_iso, hora_iso, timestamp_iso, texto_original)."""
    if not v:
        return None, None, None, None
    texto_original = str(v).strip()
    m = re.search(r'(\d{2})/(\d{2})/(\d{4})\s*-\s*(\d{2}:\d{2})', texto_original)
    if m:
        data_iso = f'{m.group(3)}-{m.group(2)}-{m.group(1)}'
        hora_iso = f'{m.group(4)}:00'
        ts_iso = f'{data_iso} {hora_iso}'
        return data_iso, hora_iso, ts_iso, texto_original
    # Se nao tiver hifen, tenta data pura
    d = _parse_date(v)
    return d, None, None, texto_original


def _parse_float(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace('R$', '').replace(' ', '').replace('.', '').replace(',', '.')
    try:
        return float(s)
    except Exception:
        return None


def _parse_int(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    try:
        return int(v)
    except Exception:
        return None


def carregar_documentos(conn, caminho_xlsx, data_extracao=None):
    """Carrega planilha de Documentos no mega.approvo_documentos com UPSERT."""
    if not data_extracao:
        data_extracao = dt.date.today().isoformat()

    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0

    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    col_idx = {h: i for i, h in enumerate(headers)}

    registros = []
    for r in rows[1:]:
        numero = _parse_int(r[col_idx.get('Número', 0)])
        if numero is None:
            continue

        tipo_doc = str(r[col_idx.get('Tipo de documento', 1)] or '').strip()
        data_doc = _parse_date(r[col_idx.get('Data', 2)])
        status = str(r[col_idx.get('Status', 3)] or '').strip()
        valor = _parse_float(r[col_idx.get('Valor (R$)', 4)])
        filial = str(r[col_idx.get('Filial', 5)] or '').strip()
        agente = str(r[col_idx.get('Agente', 6)] or '').strip()
        solicitante = str(r[col_idx.get('Solicitante', 7)] or '').strip()
        data_envio = _parse_date(r[col_idx.get('Data de envio para aprovação', 8)])
        classe_fin = str(r[col_idx.get('Classe financeira', 9)] or '').strip()
        centro_custo = str(r[col_idx.get('Centro de custo', 10)] or '').strip()
        projeto = str(r[col_idx.get('Projeto', 11)] or '').strip()
        regra_aprov = str(r[col_idx.get('Regra de aprovação', 12)] or '').strip()
        obra = _extrair_obra(filial)

        raw_data = json.dumps({
            headers[i]: r[i] for i in range(len(headers)) if i < len(r)
        }, default=str, ensure_ascii=False)

        registros.append((
            obra, numero, tipo_doc, data_doc, status, valor,
            filial, agente, solicitante, data_envio, classe_fin,
            centro_custo, projeto, regra_aprov, data_extracao, raw_data
        ))

    sql = """
        INSERT INTO mega.approvo_documentos (
          obra, numero, tipo_documento, data_documento, status, valor,
          filial, agente, solicitante, data_envio_aprovacao, classe_financeira,
          centro_custo, projeto, regra_aprovacao, data_extracao, raw_data
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (filial, tipo_documento, numero) DO UPDATE SET
          obra = EXCLUDED.obra,
          data_documento = EXCLUDED.data_documento,
          status = EXCLUDED.status,
          valor = EXCLUDED.valor,
          agente = EXCLUDED.agente,
          solicitante = EXCLUDED.solicitante,
          data_envio_aprovacao = EXCLUDED.data_envio_aprovacao,
          classe_financeira = EXCLUDED.classe_financeira,
          centro_custo = EXCLUDED.centro_custo,
          projeto = EXCLUDED.projeto,
          regra_aprovacao = EXCLUDED.regra_aprovacao,
          data_extracao = EXCLUDED.data_extracao,
          raw_data = EXCLUDED.raw_data;
    """
    cur = conn.cursor()
    cur.executemany(sql, registros)
    conn.commit()
    return len(registros)


def carregar_ocorrencias(conn, caminho_xlsx, data_extracao=None):
    """Carrega planilha de Ocorrências no mega.approvo_ocorrencias com ON CONFLICT DO NOTHING."""
    if not data_extracao:
        data_extracao = dt.date.today().isoformat()

    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0

    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    col_idx = {h: i for i, h in enumerate(headers)}

    registros = []
    # Deduplicar na memoria para evitar tentativas duplicadas no mesmo lote
    vistos = set()

    for r in rows[1:]:
        numero_doc = _parse_int(r[col_idx.get('Número do documento', 0)])
        if numero_doc is None:
            continue

        acao = str(r[col_idx.get('Ação', 1)] or '').strip()
        data_hora_raw = r[col_idx.get('Data/Hora', 2)]
        data_aprov, hora_aprov, ts, data_hora_texto = _parse_data_hora(data_hora_raw)
        aprovador = str(r[col_idx.get('Aprovador', 3)] or '').strip()
        motivo = str(r[col_idx.get('Motivo da operação', 4)] or '').strip()
        tipo_doc = str(r[col_idx.get('Tipo de documento', 5)] or '').strip()
        data_doc = _parse_date(r[col_idx.get('Data do documento', 6)])
        valor = _parse_float(r[col_idx.get('Valor (R$)', 7)])
        filial = str(r[col_idx.get('Filial', 8)] or '').strip()
        agente = str(r[col_idx.get('Agente', 9)] or '').strip()
        solicitante = str(r[col_idx.get('Solicitante', 10)] or '').strip()
        obra = _extrair_obra(filial)

        chave_unica = (filial, tipo_doc, numero_doc, data_aprov, hora_aprov, aprovador, acao)
        if chave_unica in vistos:
            continue
        vistos.add(chave_unica)

        raw_data = json.dumps({
            headers[i]: r[i] for i in range(len(headers)) if i < len(r)
        }, default=str, ensure_ascii=False)

        registros.append((
            obra, numero_doc, acao, data_aprov, hora_aprov,
            ts, data_hora_texto, aprovador, motivo, tipo_doc,
            data_doc, valor, filial, agente, solicitante, data_extracao, raw_data
        ))

    sql = """
        INSERT INTO mega.approvo_ocorrencias (
          obra, numero_documento, acao, data_aprovacao, hora_aprovacao,
          data_hora, data_hora_texto, aprovador, motivo_operacao, tipo_documento,
          data_documento, valor, filial, agente, solicitante, data_extracao, raw_data
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (filial, tipo_documento, numero_documento, data_aprovacao, hora_aprovacao, aprovador, acao)
        DO NOTHING;
    """
    cur = conn.cursor()
    cur.executemany(sql, registros)
    conn.commit()
    return len(registros)
